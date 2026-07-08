"""Render a Transcript to a court-grade DOCX that mirrors the presentation of
an official Maryland reporter's transcript: Courier New 12 pt, double-spaced,
automatic 1-per-line numbering restarting each page, a caption title page,
speaker-labeled colloquy, centered examination/section headers, event
parentheticals, and a clearly-labeled DRAFT certification page.

This is the primary output format. It deliberately mirrors the official layout
but marks the document a DRAFT / not certified.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from .model import Transcript

# Maryland circuit-court county codes (leading C-## of a case number).
_MD_COUNTIES = {
    "01": "Allegany County", "02": "Anne Arundel County", "03": "Baltimore County",
    "04": "Calvert County", "05": "Caroline County", "06": "Carroll County",
    "07": "Cecil County", "08": "Charles County", "09": "Dorchester County",
    "10": "Frederick County", "11": "Garrett County", "12": "Harford County",
    "13": "Howard County", "14": "Kent County", "15": "Montgomery County",
    "16": "Prince George's County", "17": "Queen Anne's County",
    "18": "St. Mary's County", "19": "Somerset County", "20": "Talbot County",
    "21": "Washington County", "22": "Wicomico County", "23": "Worcester County",
    "24": "Baltimore City",
}


def court_header(transcript: Transcript) -> str:
    """Full court name for the caption's top line."""
    m = re.match(r"C-(\d{2})-", transcript.case_number or "")
    if m and m.group(1) in _MD_COUNTIES:
        return f"IN THE CIRCUIT COURT FOR {_MD_COUNTIES[m.group(1)].upper()}, MARYLAND"
    if transcript.court:
        return f"IN THE {transcript.court.upper()}"
    return "TRANSCRIPT OF PROCEEDINGS"


def _apply_section(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.6)   # room for the line-number gutter
    section.right_margin = Inches(0.6)
    # Automatic line numbering: 1 per line, restart each page.
    sectPr = section._sectPr
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:start"), "1")
    ln.set(qn("w:restart"), "newPage")
    ln.set(qn("w:distance"), "360")
    # OOXML requires lnNumType before w:cols/w:docGrid in CT_SectPr.
    cols = sectPr.find(qn("w:cols"))
    if cols is not None:
        cols.addprevious(ln)
    else:
        sectPr.append(ln)
    # Different first page so the title page carries no page number.
    section.different_first_page_header_footer = True


def _add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _line(doc, text: str = "", *, align=None, bold=False, italic=False,
          left: float | None = None, first: float | None = None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if align is not None:
        p.alignment = align
    if left is not None:
        pf.left_indent = Inches(left)
    if first is not None:
        pf.first_line_indent = Inches(first)
    if text:
        r = p.add_run(text)
        r.bold = bold
        r.italic = italic
    return p


def _labels(transcript: Transcript) -> dict[str, str]:
    return {s.id: s.label for s in transcript.speakers}


def write_court_docx(transcript: Transcript, out_path) -> Path:
    out_path = Path(out_path)
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Courier New"
    normal.font.size = Pt(12)

    section = doc.sections[0]
    _apply_section(section)

    # Page number (top-right) on the non-first-page header.
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _add_page_number(hp)

    C = WD_ALIGN_PARAGRAPH.CENTER

    # ---------- Title / caption page ----------
    _line(doc, court_header(transcript), align=C)
    _line(doc)
    caption = transcript.case_name or "IN THE MATTER OF THE PROCEEDINGS"
    _line(doc, caption.upper(), align=C, bold=True)
    if transcript.case_number:
        _line(doc, f"Case No. {transcript.case_number}", align=C)
    _line(doc)
    _line(doc, "DRAFT TRANSCRIPT OF PROCEEDINGS", align=C, bold=True)
    if transcript.hearing_date:
        _line(doc, transcript.hearing_date, align=C)
    _line(doc)
    if transcript.judge:
        _line(doc, f"BEFORE:  The Honorable {transcript.judge}", left=1.0)
    _line(doc)
    # Appearances (only when speakers carry names — court transcripts).
    if any(sp.name for sp in transcript.speakers):
        _line(doc, "APPEARANCES:", left=1.0)
        for sp in transcript.speakers:
            if sp.name:
                _line(doc, f"{sp.name}  --  {sp.label}", left=1.3)
    doc.add_page_break()

    # ---------- Body: markers + segments in time order ----------
    labels = _labels(transcript)
    events = [("marker", m.time, m) for m in transcript.markers]
    events += [("segment", s.start, s) for s in transcript.segments]
    events.sort(key=lambda e: (e[1], 0 if e[0] == "marker" else 1))

    last_speaker: object = object()  # sentinel: no label emitted yet
    for kind, _t, obj in events:
        if kind == "marker":
            last_speaker = object()
            _line(doc, obj.text, align=C, bold=obj.kind == "section",
                  italic=obj.kind == "event")
        else:
            label = labels.get(obj.speaker_id, "SPEAKER")
            if obj.speaker_id != last_speaker:
                # New speaker: "LABEL:  text..." with the label indented,
                # wrapped lines returning to the text margin.
                p = _line(doc, left=0.4, first=0.4)
                p.add_run(f"{label}:  ").bold = True
                p.add_run(obj.text)
                last_speaker = obj.speaker_id
            else:
                _line(doc, obj.text, left=0.4)

    # ---------- Draft certification ----------
    doc.add_page_break()
    _line(doc, "CERTIFICATION", align=C, bold=True)
    _line(doc)
    _line(doc,
          "This is a working DRAFT transcript produced by automated speech "
          "recognition aligned to the courtroom clerk's log. It is NOT a "
          "certified transcript and has not been reviewed or certified by an "
          "approved court transcriber. It is provided for review and case "
          "preparation only.", left=0.4)

    doc.save(str(out_path))
    return out_path
