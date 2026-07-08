from docx import Document

from transcript_genie.model import Marker, Segment, Speaker, Transcript
from transcript_genie.court_docx import write_court_docx, court_header


def _transcript():
    return Transcript(
        case_name="Alan Whitfield v. Diane Whitfield",
        case_number="C-03-FM-19-001234",
        court="Court 15",
        judge="Vance",
        hearing_date="April 3, 2019",
        speakers=[
            Speaker("court", "THE COURT"),
            Speaker("def_atty", "COUNSEL FOR DEFENDANT", name="Melina Cortez"),
        ],
        segments=[
            Segment(0.0, 1.0, "Good morning.", speaker_id="court"),
            Segment(1.0, 2.0, "Still the Court.", speaker_id="court"),
            Segment(2.0, 3.0, "Good morning, Your Honor.", speaker_id="def_atty"),
        ],
        markers=[Marker(0.0, "section", "DIRECT EXAMINATION")],
    )


def test_court_header_maps_baltimore_county():
    t = _transcript()
    assert court_header(t) == "IN THE CIRCUIT COURT FOR BALTIMORE COUNTY, MARYLAND"


def test_writes_expected_content(tmp_path):
    out = write_court_docx(_transcript(), tmp_path / "c.docx")
    assert out.exists()
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "CIRCUIT COURT FOR BALTIMORE COUNTY, MARYLAND" in text
    assert "Case No. C-03-FM-19-001234" in text
    assert "DRAFT TRANSCRIPT OF PROCEEDINGS" in text
    assert "BEFORE:  The Honorable Vance" in text
    assert "APPEARANCES:" in text
    assert "Melina Cortez" in text
    assert "DIRECT EXAMINATION" in text
    assert "THE COURT:" in text
    assert text.count("THE COURT:") == 1          # consecutive dedup
    assert "NOT a certified transcript" in text


def test_has_line_numbering_and_courier(tmp_path):
    out = write_court_docx(_transcript(), tmp_path / "c.docx")
    doc = Document(str(out))
    # Courier New on the Normal style
    assert doc.styles["Normal"].font.name == "Courier New"
    # Line numbering present in the section properties
    xml = doc.sections[0]._sectPr.xml
    assert "lnNumType" in xml
    assert 'w:countBy="1"' in xml
    assert 'w:restart="newPage"' in xml
